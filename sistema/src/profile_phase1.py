"""Inspecao reproduzivel de cinco arquivos. Nao carrega PostgreSQL nem publica dados."""
import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import openpyxl

RAW = Path('data/raw')
OUT = Path('outputs/fase1')
OUT.mkdir(parents=True, exist_ok=True)

def clean(v):
    return re.sub(r'\s+', ' ', str('' if v is None else v).strip())

def profile(name, headers, rows, keys, scope, **extra):
    rows = list(rows)
    columns = []
    for i, h in enumerate(headers):
        vals = [r[i] if i < len(r) else None for r in rows]
        missing = sum(v is None or str(v).strip() == '' for v in vals)
        markers = Counter(str(v).strip() for v in vals if str(v).strip() in ('--', 'SC', '-', '...', '..', 'X'))
        types = Counter(type(v).__name__ for v in vals if v is not None)
        columns.append(dict(name=h, blank=missing, blank_pct=round(100*missing/len(rows), 3) if rows else None,
                            markers=dict(markers), storage_types=dict(types)))
    key_indices = [headers.index(k) for k in keys]
    counts = Counter(tuple(clean(r[i]) for i in key_indices) for r in rows)
    result = dict(name=name, scope=scope, rows=len(rows), column_count=len(headers), columns=columns,
                  candidate_key=keys, duplicate_excess=sum(v-1 for v in counts.values()),
                  missing_key_rows=sum(any(not clean(r[i]) for i in key_indices) for r in rows), **extra)
    (OUT/(name+'.json')).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(name, 'rows=',len(rows),'columns=',len(headers),'duplicate_excess=',result['duplicate_excess'],flush=True)
    return result

def excel(path, sheet, header_row, predicate, name, keys):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(min_row=header_row, values_only=True)
    original_headers = list(next(it))
    headers = [clean(v) for v in original_headers]
    rows = []
    scanned = 0
    for row in it:
        if not any(v is not None for v in row):
            continue
        scanned += 1
        if predicate(dict(zip(headers,row))):
            rows.append(row)
    sheets=wb.sheetnames
    wb.close()
    r = profile(name,headers,rows,keys,'Todas as linhas MT elegiveis da aba selecionada; apenas esta edicao',
                scanned_nonempty_rows=scanned, sheet=sheet, header_row=header_row, sheets=sheets,
                original_headers=original_headers)
    return headers, rows, r

def main():
    results=[]
    course_sets={}
    h, rows, r=excel(RAW/'CPC_2021.xlsx','CPC2021',1,lambda d:d.get('Sigla da UF**')=='MT',
        'cpc2021',['Código do Curso**','Ano'])
    results.append(r)
    course_sets['cpc2021']={clean(row[h.index('Código do Curso**')]) for row in rows}
    h, rows, r=excel(RAW/'conceito_enade_licenciaturas.xlsx','Conceito Enade Licenciaturas',1,
        lambda d:d.get('Sigla da UF')=='MT','enade2025',['Código do Curso','Ano'])
    field='Percentual de Concluintes Igual ou Acima do Padrão 1 de Proficiência'
    vals=[row[h.index(field)] for row in rows if isinstance(row[h.index(field)],(int,float))]
    r['proficiency_numeric_range']=[min(vals),max(vals)] if vals else None
    results.append(r)
    course_sets['enade2025']={clean(row[h.index('Código do Curso')]) for row in rows}
    with ZipFile(RAW/'microdados_censo_da_educacao_superior_2021.zip') as z:
        member=next(n for n in z.namelist() if n.upper().endswith('MICRODADOS_CADASTRO_CURSOS_2021.CSV'))
        with z.open(member) as f:
            reader=csv.reader(io.TextIOWrapper(f,encoding='cp1252'),delimiter=';')
            h=next(reader)
            selected=[]; n=0
            for row in reader:
                n+=1
                if row[h.index('CO_UF')]=='51': selected.append(row)
        keys=[k for k in ['NU_ANO_CENSO','CO_CURSO','CO_MUNICIPIO','TP_MODALIDADE_ENSINO','TP_NIVEL_ACADEMICO'] if k in h]
        r=profile('censo2021',h,selected,keys,'Todas as linhas CO_UF=51; nivel academico ainda nao filtrado',
            scanned_rows=n,member=member,encoding='cp1252',delimiter=';',uncompressed_bytes=z.getinfo(member).file_size)
        course=Counter(row[h.index('CO_CURSO')] for row in selected)
        r['course_only_duplicate_excess']=sum(v-1 for v in course.values())
        r['categories']={k:dict(Counter(row[h.index(k)] for row in selected)) for k in ['TP_NIVEL_ACADEMICO','TP_MODALIDADE_ENSINO','TP_REDE','TP_CATEGORIA_ADMINISTRATIVA'] if k in h}
        r['requested_financing_fields_present']={k:k in h for k in ['QT_ING_FIES','QT_ING_PROUNII','QT_ING_PROUNIP','QT_MAT_FIES','QT_MAT_PROUNII','QT_MAT_PROUNIP']}
        r['quantity_parse_failures']={k:sum(not re.fullmatch(r'\d+',row[i]) for row in selected) for i,k in enumerate(h) if k.startswith('QT_')}
        results.append(r)
    with ZipFile(RAW/'indicadores_trajetoria_es_2015_2024.zip') as z:
        member=next(n for n in z.namelist() if n.lower().endswith('.xlsx') and not n.startswith('~'))
        # Mantem o original intacto; extrai somente o membro XLSX selecionado.
        p=RAW/'trajetoria2015.xlsx'
        with z.open(member) as src, p.open('wb') as dst:
            import shutil
            shutil.copyfileobj(src,dst)
        h,rows,r=excel(p,'INDICADORES_TRAJETORIA',9,lambda d:str(d.get('CO_UF'))=='51',
            'trajetoria2015',['CO_CURSO','NU_ANO_INGRESSO','NU_ANO_REFERENCIA'])
        r['member']=member
        r['categories']={k:dict(Counter(str(row[h.index(k)]) for row in rows)) for k in ['TP_MODALIDADE_ENSINO','NU_ANO_REFERENCIA']}
        course_sets['trajetoria2015']={clean(row[h.index('CO_CURSO')]) for row in rows}
        results.append(r)
    data=json.loads((RAW/'mt_uf.json').read_text(encoding='utf-8'))
    flat=[]
    for var in data:
        for result in var['resultados']:
            classifications=json.dumps({c['id']: sorted(c['categoria']) for c in result['classificacoes']},sort_keys=True)
            for series in result['series']:
                loc=series['localidade']
                for year,value in series['serie'].items():
                    flat.append([var['id'],var['unidade'],loc['nivel']['id'],loc['id'],year,classifications,value])
    r=profile('ibge9514',['variavel','unidade','nivel','localidade','periodo','categorias','valor'],flat,
              ['variavel','nivel','localidade','periodo','categorias'],'Arquivo completo mt_uf.json: UF MT, nao municipios')
    r['units']=sorted(set(row[1] for row in flat))
    r['levels']=sorted(set(row[2] for row in flat))
    results.append(r)
    joins=[]
    for name in ('cpc2021','enade2025'):
        matched=len(course_sets[name] & course_sets['trajetoria2015'])
        joins.append(dict(source=name,target='trajetoria2015',source_unique_courses=len(course_sets[name]),
                          matched_courses=matched,unmatched_courses=len(course_sets[name])-matched,
                          coverage_pct=round(100*matched/len(course_sets[name]),2),
                          limitation='Cobertura por codigo apenas; nao filtra elegibilidade, SC ou horizonte analitico.'))
    (OUT/'juncoes_diagnosticas.json').write_text(json.dumps(joins,ensure_ascii=False,indent=2),encoding='utf-8')
    for r in results:
        (OUT/(r['name']+'.json')).write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'resumo.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    with (OUT/'perfil_colunas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(['fonte','coluna','linhas','vazios','percentual_vazio','marcadores','tipos_armazenados'])
        for r in results:
            for c in r['columns']:
                w.writerow([r['name'],c['name'],r['rows'],c['blank'],c['blank_pct'],json.dumps(c['markers']),json.dumps(c['storage_types'])])

if __name__=='__main__': main()
